"""T04 三产物导出回归测试（``tests/test_exporter.py``）。

覆盖架构设计第 7 节 T04 验收要点 ⑧⑨：

  * ⑧ 汇总表**恰好 13 列**、顺序固定（用 ``constants.COLUMNS`` 断言）
  * ⑨ 产物分流正确（汇总表/清单 → 成果产出；JSON → 过程产出）
  * 越界写入抛 ``OutputPathViolation``
  * CSV 使用 ``utf-8-sig``（带 BOM，防 Excel 乱码）
  * 完整 OCR 原文只进 JSON
  * 复核清单只含 ⚠️/🔵（``REVIEW_VERDICTS``）
"""

from __future__ import annotations

import json

import pytest

from core.constants import COLUMN_COUNT, COLUMNS, REVIEW_VERDICTS
from core.models import CheckResult, DeclarationRecord, OcrText, Verdict
from core.result_exporter import REVIEW_COLUMNS, ResultExporter
from infra.errors import OutputPathViolation


def make_result(
    verdict: Verdict,
    *,
    ticket: str = "SA26090215",
    part: str = "N011901-007386-001",
    order: str = "2660326M",
    brand: str = "baori",
    model: str = "A7A01G",
    detected_brand: str = "baori",
    detected_model: str = "A7A01G",
    ocr_text: str = "品牌:baori 型号:A7A01G",
    seq: int = 1,
) -> CheckResult:
    """构造一条校验结果（含完整 OCR 证据）。"""
    record = DeclarationRecord(
        ticket_no=ticket,
        part_no=part,
        order_no=order,
        product_name="FFC线材",
        seq_no=seq,
        raw_element_text=f"品牌:{brand}|型号:{model}",
        decl_brand=brand,
        decl_model=model,
    )
    ocr = OcrText(image_path=f"/img/{seq}.jpg", text_raw=ocr_text, confidence=0.95, seq=seq)
    from core.models import ImageEvidence

    record.evidences = [
        ImageEvidence(image_path=f"/img/{seq}.jpg", seq=seq, exists=True, ocr=ocr)
    ]
    return CheckResult(
        key=record.key(),
        record=record,
        verdict=verdict,
        detected_brand=detected_brand,
        detected_model=detected_model,
        evidence_text=ocr_text,
        image_paths=f"/img/{seq}.jpg",
        reason=f"{verdict.value} 测试",
        evidence_images=record.evidences,
    )


@pytest.fixture
def exporter(temp_project) -> ResultExporter:
    """绑定临时工程空间的导出器。"""
    return ResultExporter(
        result_dir=temp_project["result"],
        process_dir=temp_project["process"],
    )


# ══════════════════════════════════════════════════════════════════
#  ⑧ 汇总表 13 列
# ══════════════════════════════════════════════════════════════════


class TestSummaryColumns:
    """汇总表恰好 13 列，顺序固定。"""

    def test_columns_count(self) -> None:
        assert len(COLUMNS) == COLUMN_COUNT == 13

    def test_exporter_columns_is_constants(self) -> None:
        """``ResultExporter.COLUMNS`` 是 ``constants.COLUMNS`` 的唯一事实来源。"""
        assert ResultExporter.COLUMNS == COLUMNS

    def test_export_13_columns(self, exporter: ResultExporter) -> None:
        """导出的 xlsx 表头恰好 13 列且顺序一致。"""
        from openpyxl import load_workbook

        results = [
            make_result(Verdict.PASS),
            make_result(Verdict.FAIL, seq=2),
        ]
        path = exporter.export_summary(results, "SA26090215")
        assert path.is_file()

        workbook = load_workbook(str(path), read_only=True)
        sheet = workbook.active
        header = [cell for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        assert len(header) == 13
        assert list(header) == COLUMNS
        workbook.close()

    def test_summary_row_count(self, exporter: ResultExporter) -> None:
        """表头 1 行 + 数据 N 行。"""
        from openpyxl import load_workbook

        results = [make_result(Verdict.PASS, seq=i) for i in range(1, 6)]
        path = exporter.export_summary(results, "SA26090215")
        workbook = load_workbook(str(path), read_only=True)
        sheet = workbook.active
        assert sheet.max_row == 6
        workbook.close()

    def test_to_row_length(self) -> None:
        """``CheckResult.to_row()`` 恰好 13 个值。"""
        assert len(make_result(Verdict.PASS).to_row()) == 13

    def test_summary_path_name(self, exporter: ResultExporter) -> None:
        path = exporter.export_summary([], "SA26090215")
        assert path.name == "校验汇总表_SA26090215.xlsx"
        assert path.parent.name == "成果产出"


# ══════════════════════════════════════════════════════════════════
#  v0.3.6 · 第 10 列「判定依据」（用户裁定 2026-09-17）
#
#  原「差异备注」**只在存在差异时才有内容** → ✅ 合格行整格为空，合格记录拿不到依据。
#  现改为「判定依据」：对**每一条记录**都写明依据 ——
#  ``reason``（判定依据句）＋ 差异明细（含逐字符差异）＋ 复核备注。
#  ⚠️ 改的是**列名与内容口径**，列数与顺序**不变**（仍恰好 13 列）。
# ══════════════════════════════════════════════════════════════════

#: 第 10 列（0-based = 9）在 ``to_row()`` 里的下标
BASIS_COLUMN_INDEX = 9


class TestVerdictBasisColumn:
    """第 10 列列名与内容口径：**四类结论都必须落盘判定依据**。"""

    def test_column_renamed(self) -> None:
        """列名「差异备注」→「判定依据」，且列数与顺序不变。"""
        assert COLUMNS[BASIS_COLUMN_INDEX] == "判定依据"
        assert "差异备注" not in COLUMNS
        assert len(COLUMNS) == COLUMN_COUNT == 13

    @pytest.mark.parametrize("verdict", list(Verdict), ids=lambda v: v.name)
    def test_every_verdict_has_basis(self, verdict: Verdict) -> None:
        """✅/❌/⚠️/🔵 **四类结论**在第 10 列都必须有判定依据（不得为空）。"""
        result = make_result(verdict)
        basis = result.to_row()[BASIS_COLUMN_INDEX]
        assert basis, f"{verdict.value} 的判定依据为空 —— 又退回「只有差异才有内容」"
        assert basis == result.reason

    def test_diff_details_appended_after_reason(self) -> None:
        """有差异时：reason 在前，差异明细（含逐字符差异）在后。"""
        from core.models import DifferenceDetail

        result = make_result(Verdict.FAIL)
        result.reason = "品牌：申报『baori』与图片『Daewoo』明确不一致"
        result.differences = [
            DifferenceDetail(
                field="品牌",
                declared_value="baori",
                detected_value="Daewoo",
                char_diffs=["『申报b』→『图片D』（第0位起）"],
                note="外箱整机品牌，不构成本体证据",
            )
        ]
        basis = result.to_row()[BASIS_COLUMN_INDEX]
        assert basis.startswith("品牌：申报『baori』与图片『Daewoo』明确不一致")
        assert "差异点：" in basis
        assert "外箱整机品牌" in basis

    def test_note_deduped_against_reason(self) -> None:
        """差异明细的 note 若已出现在 reason 里 → **不再重复**。

        ``judge_engine`` 会把字段级依据**同时**写进 ``reason`` 与
        ``DifferenceDetail.note``；两段拼成一格时不去重就会出现同句两三遍。
        """
        from core.models import DifferenceDetail

        note = "品牌：申报『baori』与图片『NMY』明确不一致"
        result = make_result(Verdict.FAIL)
        result.reason = note
        result.differences = [
            DifferenceDetail(
                field="品牌",
                declared_value="baori",
                detected_value="NMY",
                char_diffs=["『申报baori』→『图片NMY』（第0位起）"],
                note=note,
            )
        ]
        basis = result.to_row()[BASIS_COLUMN_INDEX]
        assert basis.count(note) == 1, f"同句重复：{basis!r}"
        assert "差异点：" in basis

    def test_note_kept_when_not_in_reason(self) -> None:
        """note 不在 reason 里时**必须保留**（如「外箱整机品牌，不构成本体证据」）。"""
        from core.models import DifferenceDetail

        result = make_result(Verdict.FAIL)
        result.reason = "品牌：申报『baori』与图片『Daewoo』明确不一致"
        result.differences = [
            DifferenceDetail(
                field="品牌",
                declared_value="baori",
                detected_value="Daewoo",
                note="外箱整机品牌，不构成本体证据",
            )
        ]
        basis = result.to_row()[BASIS_COLUMN_INDEX]
        assert "外箱整机品牌，不构成本体证据" in basis

    def test_reviewer_note_appended(self) -> None:
        """复核备注一并写进「判定依据」。"""
        result = make_result(Verdict.NO_MARK)
        result.reason = "图片内未识别到品牌/型号文字，转人工复核"
        result.reviewer_note = "已看图确认标签无品牌"
        basis = result.to_row()[BASIS_COLUMN_INDEX]
        assert "复核备注：已看图确认标签无品牌" in basis

    def test_exported_sheet_basis_not_empty_for_pass(
        self, exporter: ResultExporter
    ) -> None:
        """端到端：导出的 xlsx 里，✅ 行的第 10 列也必须有依据。"""
        from openpyxl import load_workbook

        path = exporter.export_summary([make_result(Verdict.PASS)], "SA26090215")
        workbook = load_workbook(str(path), read_only=True)
        sheet = workbook.active
        header = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        assert len(header) == 13
        assert header[BASIS_COLUMN_INDEX] == "判定依据"
        row = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
        assert row[BASIS_COLUMN_INDEX], "✅ 行的判定依据为空"
        assert row[BASIS_COLUMN_INDEX] == row[12]  # 无差异时 = reason（= 第 13 列）
        workbook.close()


# ══════════════════════════════════════════════════════════════════
#  ⑨ 产物分流
# ══════════════════════════════════════════════════════════════════


class TestProductRouting:
    """汇总表/清单 → 成果产出；JSON → 过程产出。"""

    def test_summary_in_result_dir(self, exporter: ResultExporter, temp_project) -> None:
        path = exporter.export_summary([make_result(Verdict.PASS)], "SA26090215")
        assert path.parent == temp_project["result"]

    def test_review_csv_in_result_dir(self, exporter: ResultExporter, temp_project) -> None:
        path = exporter.export_review_csv([make_result(Verdict.NO_MARK)], "SA26090215")
        assert path.parent == temp_project["result"]

    def test_detail_json_in_process_dir(self, exporter: ResultExporter, temp_project) -> None:
        path = exporter.export_detail_json([make_result(Verdict.PASS)], "SA26090215")
        assert path.parent == temp_project["process"]
        assert path.name == "校验详细日志_SA26090215_累积.json"

    def test_export_all(self, exporter: ResultExporter, temp_project) -> None:
        """``export_all`` 一次性产出三产物且分流正确。"""
        results = [make_result(Verdict.PASS), make_result(Verdict.NO_MARK, seq=2)]
        paths = exporter.export_all(results, "SA26090215")
        assert paths["summary"].parent == temp_project["result"]
        assert paths["review"].parent == temp_project["result"]
        assert paths["detail"].parent == temp_project["process"]
        for path in paths.values():
            assert path.is_file()

    def test_json_not_in_result_dir(self, exporter: ResultExporter) -> None:
        """详细日志 JSON 绝不落成果产出。"""
        path = exporter.export_detail_json([make_result(Verdict.PASS)], "SA26090215")
        assert path.parent.name == "过程产出"


# ══════════════════════════════════════════════════════════════════
#  越界写入 → OutputPathViolation
# ══════════════════════════════════════════════════════════════════


class TestPathViolation:
    """产物分流铁律：越界写入抛 ``OutputPathViolation``。"""

    def test_summary_out_of_root(self, temp_project) -> None:
        """汇总表写入越界（指定到输入目录）→ 抛异常。"""
        exporter = ResultExporter(
            result_dir=temp_project["result"],
            process_dir=temp_project["process"],
        )
        with pytest.raises(OutputPathViolation):
            exporter.export_summary(
                [make_result(Verdict.PASS)],
                "SA26090215",
                out_dir=temp_project["input"],
            )

    def test_detail_out_of_root(self, temp_project) -> None:
        """JSON 写入成果产出 → 抛异常。"""
        exporter = ResultExporter(
            result_dir=temp_project["result"],
            process_dir=temp_project["process"],
        )
        with pytest.raises(OutputPathViolation):
            exporter.export_detail_json(
                [make_result(Verdict.PASS)],
                "SA26090215",
                proc_dir=temp_project["result"],
            )

    def test_assert_helpers(self, temp_project) -> None:
        """公开断言辅助。"""
        ResultExporter.assert_target_in_result(
            temp_project["result"] / "x.xlsx", temp_project["result"]
        )
        ResultExporter.assert_target_in_process(
            temp_project["process"] / "x.json", temp_project["process"]
        )
        with pytest.raises(OutputPathViolation):
            ResultExporter.assert_target_in_result(
                temp_project["input"] / "x.xlsx", temp_project["result"]
            )


# ══════════════════════════════════════════════════════════════════
#  复核清单
# ══════════════════════════════════════════════════════════════════


class TestReviewCsv:
    """待人工复核清单。"""

    def test_utf8_bom(self, exporter: ResultExporter, temp_project) -> None:
        """CSV 必须带 BOM（utf-8-sig）。"""
        path = exporter.export_review_csv([make_result(Verdict.NO_MARK)], "SA26090215")
        raw = path.read_bytes()
        assert raw.startswith(b"\xef\xbb\xbf"), "CSV 缺少 UTF-8 BOM，Excel 会乱码"

    def test_only_review_verdicts(self, exporter: ResultExporter) -> None:
        """清单只含 ⚠️ + 🔵，不含 ✅/❌。"""
        import csv

        results = [
            make_result(Verdict.PASS, seq=1),
            make_result(Verdict.FAIL, seq=2),
            make_result(Verdict.NO_MARK, seq=3),
            make_result(Verdict.NO_IMAGE, seq=4),
        ]
        path = exporter.export_review_csv(results, "SA26090215")
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))

        assert rows[0] == REVIEW_COLUMNS
        data_rows = rows[1:]
        assert len(data_rows) == 2, "复核清单应只含 ⚠️ + 🔵 两条"

    def test_review_verdicts_constant(self) -> None:
        """``REVIEW_VERDICTS`` 只含 ⚠️ + 🔵。"""
        assert set(REVIEW_VERDICTS) == {Verdict.NO_MARK, Verdict.NO_IMAGE}

    def test_problem_column_shares_basis_exit(
        self, exporter: ResultExporter
    ) -> None:
        """「问题说明」与汇总表第 10 列**同源同文**（v0.3.6 单一出口）。

        此前两处各写一遍 ``reason + d.summary()``，同一条依据（既在 ``reason``
        又在 ``DifferenceDetail.note``）会拼出同句两三遍。收口到
        ``CheckResult.verdict_basis()`` 后，两边必须逐字相同。
        """
        import csv

        from core.models import DifferenceDetail

        note = "品牌：申报『baori』与图片『NMY』明确不一致"
        result = make_result(Verdict.NO_MARK)
        result.reason = note
        result.differences = [
            DifferenceDetail(
                field="品牌",
                declared_value="baori",
                detected_value="NMY",
                char_diffs=["『申报baori』→『图片NMY』（第0位起）"],
                note=note,
            )
        ]

        path = exporter.export_review_csv([result], "SA26090215")
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
        problem = rows[1][REVIEW_COLUMNS.index("问题说明")]

        basis = result.to_row()[BASIS_COLUMN_INDEX]
        assert problem == basis
        assert problem.count(note) == 1, f"同句重复：{problem!r}"

    def test_csv_path_name(self, exporter: ResultExporter) -> None:
        path = exporter.export_review_csv([], "SA26090215")
        assert path.name == "SA26090215_待人工复核清单.csv"

    def test_problem_column_no_tuple_bug(self, exporter: ResultExporter) -> None:
        """R10：问题说明列不得因元组取值 bug 以'零'开头。"""
        import csv

        results = [make_result(Verdict.NO_MARK)]
        path = exporter.export_review_csv(results, "SA26090215")
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
        header = rows[0]
        problem_idx = header.index("问题说明")
        problem = rows[1][problem_idx]
        assert not problem.startswith("0"), "问题说明疑似元组取值 bug"


# ══════════════════════════════════════════════════════════════════
#  详细日志 JSON
# ══════════════════════════════════════════════════════════════════


class TestDetailJson:
    """详细日志 JSON（含完整 OCR 原文）。"""

    def test_contains_full_ocr(self, exporter: ResultExporter) -> None:
        """JSON 含完整 OCR 原文（这是唯一允许存全文的产物）。"""
        long_ocr = "品牌:baori 型号:A7A01G " + "X" * 300
        result = make_result(Verdict.PASS, ocr_text=long_ocr)
        path = exporter.export_detail_json([result], "SA26090215")
        payload = json.loads(path.read_text(encoding="utf-8"))
        # 定位证据里的 text_raw
        text_raw = payload["results"][0]["evidence_images"][0]["ocr"]["text_raw"]
        assert text_raw == long_ocr

    def test_counts_present(self, exporter: ResultExporter) -> None:
        """JSON 含四类判定计数。"""
        results = [
            make_result(Verdict.PASS, seq=1),
            make_result(Verdict.FAIL, seq=2),
            make_result(Verdict.NO_MARK, seq=3),
            make_result(Verdict.NO_IMAGE, seq=4),
        ]
        path = exporter.export_detail_json(results, "SA26090215")
        payload = json.loads(path.read_text(encoding="utf-8"))
        assert payload["total"] == 4
        assert payload["counts"]["PASS"] == 1
        assert payload["counts"]["FAIL"] == 1
        assert payload["counts"]["NO_MARK"] == 1
        assert payload["counts"]["NO_IMAGE"] == 1

    def test_extra_meta(self, exporter: ResultExporter) -> None:
        path = exporter.export_detail_json(
            [make_result(Verdict.PASS)],
            "SA26090215",
            extra={"fingerprint": {"ticket_no": "SA26090215"}},
        )
        payload = json.loads(path.read_text(encoding="utf-8"))
        assert payload["meta"]["fingerprint"]["ticket_no"] == "SA26090215"

    def test_schema(self, exporter: ResultExporter) -> None:
        path = exporter.export_detail_json([make_result(Verdict.PASS)], "SA26090215")
        payload = json.loads(path.read_text(encoding="utf-8"))
        assert payload["schema"] == 2
        assert payload["ticket_no"] == "SA26090215"


# ══════════════════════════════════════════════════════════════════
#  中文路径
# ══════════════════════════════════════════════════════════════════


class TestChinesePath:
    """中文路径安全写入。"""

    def test_chinese_result_dir(self, tmp_path) -> None:
        result_dir = tmp_path / "成果产出测试"
        process_dir = tmp_path / "过程产出测试"
        exporter = ResultExporter(result_dir=result_dir, process_dir=process_dir)
        summary = exporter.export_summary([make_result(Verdict.PASS)], "SA26090215")
        detail = exporter.export_detail_json([make_result(Verdict.PASS)], "SA26090215")
        assert summary.is_file()
        assert detail.is_file()
        assert "成果产出" in str(summary)

    def test_input_not_modified(self, temp_project) -> None:
        """导出不得触碰原始输入目录。"""
        input_dir = temp_project["input"]
        sentinel = input_dir / "sentinel.txt"
        sentinel.write_text("do not touch", encoding="utf-8")
        before = sentinel.stat().st_mtime_ns

        exporter = ResultExporter(
            result_dir=temp_project["result"], process_dir=temp_project["process"]
        )
        exporter.export_all([make_result(Verdict.PASS)], "SA26090215")

        assert sentinel.read_text(encoding="utf-8") == "do not touch"
        assert sentinel.stat().st_mtime_ns == before
